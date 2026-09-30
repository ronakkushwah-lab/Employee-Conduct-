# from groups_manager.models import Group,GroupType, Member
from account.forms import SignUpForm
from account.models import User, CompanyStaff, Company
from django.views.generic import TemplateView, CreateView
from django.contrib.auth.models import Group, Permission
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import ensure_csrf_cookie, csrf_protect, csrf_exempt
from employee.models import Department, Designation, Employee, Attendance
from managers.models import Manager
from django.utils import timezone
from django.conf import settings
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http.response import HttpResponseRedirect, JsonResponse
from django.shortcuts import render, HttpResponse, redirect, get_object_or_404
from django.urls import reverse_lazy, reverse
# from django.contrib.auth.models import Group, User
from django.views.generic import View, TemplateView, UpdateView
from django.db import IntegrityError
from django.contrib import messages
from django.contrib.contenttypes.models import ContentType
from django.contrib.auth.hashers import make_password
from django.contrib.auth.hashers import check_password
import sweetify
from django.core.mail import EmailMessage
import random


# Signs Up View

#
class SignUpView(CreateView):
    form_class = SignUpForm
    success_url = reverse_lazy('signin')
    template_name = 'account/signup.html'


@method_decorator(csrf_exempt, name='dispatch')
class SignInView(View):
    def post(self, request):
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '')
        if not email:
            from biometric.views import iclock_cdata
            return iclock_cdata(request)
        user = authenticate(username=email, password=password)
        if user is not None:
            if user.is_active:
                login(request, user)
                if user.groups.all().exists() or user.is_company_admin:
                    return HttpResponseRedirect('/administration/index')

                if user.is_staff:
                    return HttpResponseRedirect('/managers/dashboard')

                if user.is_employee:
                    return HttpResponseRedirect("/employee/employee_dashboard")

                else:
                    return HttpResponseRedirect(settings.LOGIN_URL)
            else:
                return HttpResponse("Inactive user.")
        else:

            return HttpResponseRedirect(settings.LOGIN_URL)

    def get(self, request):
        return render(request, "account/login.html")


class LogoutView(View):
    def get(self, request):
        """
        Fully log the user out of the custom CompanyStaff session.
        After this, browser back button should not restore an active session.
        """
        # If we are using the custom CompanyStaff-based auth, mark it as logged out
        company_staff_id = request.session.get('company_staff_id')
        if company_staff_id:
            try:
                staff = CompanyStaff.objects.get(pk=company_staff_id)
                staff.is_authenticated = False
                staff.save()
            except CompanyStaff.DoesNotExist:
                pass

        # Clear all session data so no stale login state remains
        request.session.flush()

        # Also log out of Django's auth system (for safety, if ever used)
        logout(request)

        return HttpResponseRedirect(settings.LOGIN_URL)


# Role
@method_decorator(user_passes_test(lambda u: u.is_superuser), name='post')
class RegisterRole(View):
    def post(self, request):
        role_name = request.POST['role']
        try:
            group = Group.objects.create(name=role_name)
            sweetify.success(self.request, f'{group} is created', button='Ok', timer=3000)
        except IntegrityError as e:
            sweetify.success(self.request, f"{group} is Already exist, button='Ok'", timer=3000)
        groups = Group.objects.all()
        return render(request, "account/role.html", {'groups': groups})

    def get(self, request):
        groups = Group.objects.all()
        return render(request, "account/role.html", {'groups': groups})


class RemoveRole(View):
    def get(self, request, name):
        try:

            groups = Group.objects.get(name=name)

            if User.objects.filter(groups__name=groups).exists():
                messages.error(request,
                               f'Cant delete {groups} ,Delete assigned Users  First and Try Again! <a href="/usertorole/{groups}"> click Here </a>',
                               extra_tags='safe')

            else:
                groups.delete()
                # messages.success(request,f"{groups} Deleted Successfully")
                sweetify.success(self.request, f'{groups} is Deleted', button='Ok', timer=3000)

        except Group.DoesNotExist:
            messages.error(request, "Role already Deleted or Not Created")
        return HttpResponseRedirect('/role')


class RemoveUserToRole(View):
    def get(self, request, name, id):
        role = Group.objects.get(name=name)
        user = User.objects.get(id=id)
        user.is_admin = False
        user.save()
        user.groups.remove(role)
        # messages.warning(request,f"{user} is removed from {role} ") 
        sweetify.info(self.request, f"{user} is removed from {role} ", button='Ok', timer=3000)
        return redirect('/usertorole/' + str(role))


class demoview(TemplateView):
    template_name = "account/demo.html"


class UserToRole(View):
    def get(self, request, name):
        role = Group.objects.get(name=name)
        employees = Employee.objects.all()
        role_user = User.objects.filter(groups__name=role)
        for user in role_user:
            user.is_admin = True
            user.save()
        return render(request, 'account/usertorole.html',
                      {'role': role, 'employees': employees, 'role_user': role_user
                       })

    def post(self, request, name):
        role = Group.objects.get(name=name)
        employe = request.POST['employee']
        user = User.objects.get(email=employe)
        userIngroup = user.groups.all().exists()

        if userIngroup != True:
            user.groups.add(role)
            # messages.info(request,f"Congratulation {user} become a {role} ")
            sweetify.success(self.request, f"Congratulation {user} become a {role} ", button='Ok', timer=3000)
        else:
            userInWhichgroup = user.groups.all()
            for userRole in userInWhichgroup:
                messages.warning(request, f"Sorry {user} is Already having {userRole} Role ")
        return redirect('/usertorole/' + str(role))


class RolePermissionView(View):
    def get(self, request, name):
        role = Group.objects.get(name=name)
        permissions = role.permissions.all()

        for permission in permissions:
            if permission.codename == 'view_employee':
                view_employee = 'True'
            else:
                view_employee = 'False'

            if permission.codename == 'add_employee':
                add_employee = 'True'
            else:
                add_employee = 'False'

            if permission.codename == 'change_employee':
                change_employee = 'True'
            else:
                change_employee = 'False'
            if permission.codename == 'delete_employee':
                delete_employee = 'True'
            else:
                delete_employee = 'False'
        # ______________employee end_________________________________________
        return render(request, 'account/add_roles_permission.html',
                      {'role': role,
                       # 'add_employee':add_employee,
                       # 'view_employee':view_employee,
                       # 'change_employee':change_employee,
                       # 'delete_employee':delete_employee

                       })

    def post(self, request, name):
        role = Group.objects.get(name=name)
        view_employee = request.POST['view_employee']
        add_employee = request.POST['add_employee']
        change_employee = request.POST['change_employee']
        delete_employee = request.POST['delete_employee']

        content_type = ContentType.objects.get_for_model(Employee, for_concrete_model=False)
        employee_permision = Permission.objects.filter(content_type=content_type)
        for permission in employee_permision:
            if permission.codename == 'view_employee':
                if view_employee == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'add_employee':
                if add_employee == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'change_employee':
                if change_employee == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'delete_employee':
                if delete_employee == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
        sweetify.info(self.request, 'Permision Granted', button='Ok', timer=3000)

        view_department = request.POST['view_department']
        add_department = request.POST['add_department']
        change_department = request.POST['change_department']
        delete_department = request.POST['delete_department']

        content_type = ContentType.objects.get_for_model(Department, for_concrete_model=False)
        department_permision = Permission.objects.filter(content_type=content_type)
        for permission in department_permision:
            if permission.codename == 'view_department':
                if view_department == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'add_department':
                if add_department == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'change_department':
                if change_department == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'delete_department':
                if delete_department == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
        sweetify.info(self.request, 'Permision Granted', button='Ok', timer=3000)

        view_designation = request.POST['view_designation']
        add_designation = request.POST['add_designation']
        change_designation = request.POST['change_designation']
        delete_designation = request.POST['delete_designation']

        content_type = ContentType.objects.get_for_model(Designation, for_concrete_model=False)
        designation_permision = Permission.objects.filter(content_type=content_type)
        for permission in designation_permision:
            if permission.codename == 'view_designation':
                if view_designation == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'add_designation':
                if add_designation == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'change_designation':
                if change_designation == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'delete_designation':
                if delete_designation == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
        sweetify.info(self.request, 'Permision Granted', button='Ok', timer=3000)

        view_goal = request.POST['view_goal']
        add_goal = request.POST['add_goal']
        change_goal = request.POST['change_goal']
        delete_goal = request.POST['delete_goal']

        content_type = ContentType.objects.get_for_model(Goal, for_concrete_model=False)
        goal_permision = Permission.objects.filter(content_type=content_type)
        for permission in goal_permision:
            if permission.codename == 'view_goal':
                if view_goal == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'add_goal':
                if add_goal == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'change_goal':
                if change_goal == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
            if permission.codename == 'delete_goal':
                if delete_goal == 'True':
                    role.permissions.add(permission)
                else:
                    role.permissions.remove(permission)
        sweetify.info(self.request, 'Permision Granted', button='Ok', timer=3000)
        return render(request, 'account/add_roles_permission.html',
                      {'role': role,
                       'view_employee': view_employee,
                       'add_employee': add_employee,
                       'change_employee': change_employee,
                       'delete_employee': delete_employee,

                       'view_department': view_department,
                       'add_department': add_department,
                       'change_department': change_department,
                       'delete_department': delete_department,

                       'view_designation': view_designation,
                       'add_designation': add_designation,
                       'change_designation': change_designation,
                       'delete_designation': delete_designation,

                       'view_goal': view_goal,
                       'add_goal': add_goal,
                       'change_goal': change_goal,
                       'delete_goal': delete_goal,
                       })


def signup(request):
    messages.error(request, 'Public registration is disabled. Please contact your company administrator to obtain login credentials.')
    return redirect('signin')


@method_decorator(csrf_exempt, name='dispatch')
class Login(View):
    return_url = None

    def post(self, request):
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '')
        if not email:
            from biometric.views import iclock_cdata
            return iclock_cdata(request)
        company_staff = CompanyStaff.get_Staff_by_email(email)
        try:
            if company_staff:
                request.session["new_notification"] = getattr(company_staff, 'new_notification', False)
                if company_staff.is_active:
                    flag = check_password(password, company_staff.password)
                    if not flag:
                        messages.info(request, "Incorrect Email or Password")
                        return HttpResponseRedirect('/hrms/')

                    # Check 60-day (2-month) password expiration policy
                    if company_staff.is_password_expired(expiry_days=60):
                        request.session['company_staff_id'] = company_staff.id
                        request.session['password_expired'] = True
                        return redirect('expired_password_change')

                    # Password is valid and within 60 days
                    if 'password_expired' in request.session:
                        del request.session['password_expired']
                    company_staff.is_authenticated = True
                    company_staff.save()
                    request.session['company_staff_id'] = company_staff.id

                    role = getattr(company_staff, 'role', None) or self._role_from_flags(company_staff)
                    is_hr_user = role == CompanyStaff.ROLE_HR or getattr(company_staff, 'is_hr', False)

                    # Smart auto-routing based on role - redirect directly to portal dashboards
                    if role == CompanyStaff.ROLE_SUPERADMIN:
                        return HttpResponseRedirect('/hrms/superadmin/')
                    if (role == CompanyStaff.ROLE_ADMIN or company_staff.is_company_admin) and company_staff.company_id:
                        return HttpResponseRedirect(f"/hrms/dashboard/admin/{company_staff.company_id}/{company_staff.pk}/")
                    if is_hr_user and company_staff.company_id:
                        return HttpResponseRedirect(f"/hrms/dashboard/hr/{company_staff.company_id}/{company_staff.pk}/")
                    if (role == CompanyStaff.ROLE_MANAGER or company_staff.is_manager) and company_staff.company_id:
                        return HttpResponseRedirect(f"/hrms/managers/dashboard/{company_staff.company_id}/{company_staff.pk}/")
                    if (role == CompanyStaff.ROLE_EMPLOYEE or company_staff.is_employee) and company_staff.company_id:
                        return HttpResponseRedirect(f"/hrms/employee/employee_dashboard/{company_staff.company_id}/{company_staff.pk}/")

                    # Fallbacks
                    if company_staff.is_company_admin and company_staff.company_id:
                        return HttpResponseRedirect(f"/hrms/administration/index/{company_staff.company_id}/{company_staff.pk}/")
                    if company_staff.is_manager and company_staff.company_id:
                        return HttpResponseRedirect(f"/hrms/managers/dashboard/{company_staff.company_id}/{company_staff.pk}/")
                    if company_staff.is_employee and company_staff.company_id:
                        return HttpResponseRedirect(f"/hrms/employee/employee_dashboard/{company_staff.company_id}/{company_staff.pk}/")
                    return HttpResponseRedirect('/hrms/')
                else:
                    messages.error(request, "Your account is inactive. Please contact your administrator.")
                    return HttpResponseRedirect('/hrms/')
            else:
                messages.info(request, "Incorrect Email or Password")
                return HttpResponseRedirect('/hrms/')
        except Exception as e:
            messages.error(request, f"Login error: {str(e)}")
            return HttpResponseRedirect('/hrms/')

    @staticmethod
    def _role_from_flags(company_staff):
        """Fallback: derive role from legacy flags if role field is empty."""
        if getattr(company_staff, 'is_hr', False):
            return CompanyStaff.ROLE_HR
        if company_staff.is_company_admin:
            return CompanyStaff.ROLE_ADMIN
        if company_staff.is_manager:
            return CompanyStaff.ROLE_MANAGER
        if company_staff.is_employee:
            return CompanyStaff.ROLE_EMPLOYEE
        return CompanyStaff.ROLE_EMPLOYEE


    def get(self, request):
        # Flush any stale workflow messages from session so they never appear on login page
        from django.contrib.messages import get_messages
        storage = get_messages(request)
        for _ in storage:
            pass
        return render(request, "account/login.html")


def superadmin_dashboard(request):
    """Dashboard for superadmin role - redirect directly to superadmin panel."""
    return HttpResponseRedirect('/hrms/superadmin/')


def admin_dashboard(request, company_id, company_staff_id):
    """Dashboard for admin role - load administration portal directly under admin dashboard route."""
    return hr_dashboard(request, company_id, company_staff_id)


def manager_dashboard(request, company_id, company_staff_id):
    """Dashboard for manager role - redirect directly to manager portal."""
    return HttpResponseRedirect(f'/hrms/managers/dashboard/{company_id}/{company_staff_id}/')


def employee_dashboard(request, company_id, company_staff_id):
    """Dashboard for employee role - redirect directly to employee panel."""
    return HttpResponseRedirect(f'/hrms/employee/employee_dashboard/{company_id}/{company_staff_id}/')


def forgotpass(request):
    context = {}
    if request.method == "POST":
        email = request.POST["email"]
        password = request.POST["password"]

        user = get_object_or_404(CompanyStaff, email=email)
        user.password = make_password(password)
        user.password_changed_at = timezone.now()
        user.save()

        from account.models import User as AccountUser
        dj_user = AccountUser.objects.filter(email=user.email).first()
        if dj_user:
            dj_user.set_password(password)
            dj_user.save()

        return HttpResponseRedirect('/hrms/')

    return render(request, "account/forgot_pass.html", context)


from django.contrib.auth.tokens import PasswordResetTokenGenerator
from django.utils.http import urlsafe_base64_encode, urlsafe_base64_decode
from django.utils.encoding import force_bytes, force_str

class StaffPasswordResetTokenGenerator(PasswordResetTokenGenerator):
    def _make_hash_value(self, user, timestamp):
        email = user.email or ''
        password = user.password or ''
        return f"{user.pk}{password}{timestamp}{email}"

staff_token_generator = StaffPasswordResetTokenGenerator()


def reset_password(request):
    import logging
    logger = logging.getLogger(__name__)
    email_address = request.GET.get("email", "").strip()
    if not email_address:
        return JsonResponse({"status": "failed"})
    try:
        user = get_object_or_404(CompanyStaff, email=email_address)
        
        uidb64 = urlsafe_base64_encode(force_bytes(user.pk))
        token = staff_token_generator.make_token(user)
        reset_path = reverse('reset_password_confirm_staff', kwargs={'uidb64': uidb64, 'token': token})
        reset_link = request.build_absolute_uri(reset_path)

        msz = (
            f"Dear {user.email},\n\n"
            f"You requested a password reset for your HRMS account.\n\n"
            f"Click the link below to reset your password:\n"
            f"{reset_link}\n\n"
            f"This link is valid for one-time use only.\n"
            f"If you did not request this password reset, please ignore this email.\n\n"
            f"Thanks & Regards,\nHRMS Portal"
        )
        from_email = (getattr(settings, 'DEFAULT_FROM_EMAIL', None) or getattr(settings, 'EMAIL_HOST_USER', None) or 'noreply@eagleincloud.io').strip()
        if not from_email:
            from_email = 'noreply@eagleincloud.io'
        try:
            msg = EmailMessage(
                subject="Password Reset - HRMS Portal",
                body=msz,
                from_email=from_email,
                to=[user.email],
            )
            msg.send(fail_silently=False)
            logger.info("Password reset email successfully sent to %s", user.email)
        except Exception as e:
            logger.exception("Failed to send reset email to %s: %s", user.email, e)
            logger.info("Reset Link generated for %s: %s", user.email, reset_link)

        # Secure Option 1: Never expose reset token or link in public HTTP response
        return JsonResponse({"status": "sent", "email": user.email})
    except Exception:
        # Generic response to prevent user enumeration
        return JsonResponse({"status": "sent", "email": email_address})


@csrf_exempt
def reset_password_confirm(request, uidb64, token):
    try:
        uid = force_str(urlsafe_base64_decode(uidb64))
        user = CompanyStaff.objects.get(pk=uid)
    except (TypeError, ValueError, OverflowError, CompanyStaff.DoesNotExist):
        user = None

    if user is not None and staff_token_generator.check_token(user, token):
        if request.method == 'POST':
            password = request.POST.get('password')
            confirm_password = request.POST.get('confirm_password')
            if not password or password != confirm_password:
                return render(request, 'account/password_reset_confirm.html', {
                    'validlink': True,
                    'error': 'Passwords do not match. Please try again.',
                    'email': user.email
                })
            user.password = make_password(password)
            user.password_changed_at = timezone.now()
            user.save()

            from account.models import User as AccountUser
            dj_user = AccountUser.objects.filter(email=user.email).first()
            if dj_user:
                dj_user.set_password(password)
                dj_user.save()

            sweetify.success(request, 'Password Reset Successful', text='Your password has been changed. Please sign in.', persistent='OK')
            messages.success(request, "Your password has been reset successfully! Please sign in with your new password.")
            return redirect('/hrms/')

        return render(request, 'account/password_reset_confirm.html', {'validlink': True, 'email': user.email})
    else:
        return render(request, 'account/password_reset_confirm.html', {'validlink': False})


def expired_password_change(request):
    """
    Forces user to change password after 60 days (2-month expiration policy).
    Once updated, password_changed_at is reset to now and user is routed to their dashboard.
    """
    company_staff_id = request.session.get('company_staff_id')
    if not company_staff_id:
        return redirect('/hrms/')

    staff = get_object_or_404(CompanyStaff, id=company_staff_id)

    if request.method == 'POST':
        old_password = request.POST.get('old_password', '').strip()
        new_password = request.POST.get('new_password', '').strip()
        confirm_password = request.POST.get('confirm_password', '').strip()

        # 1. Validate old password
        if not check_password(old_password, staff.password):
            messages.error(request, 'Current password entered is incorrect.')
            return render(request, 'account/expired_password.html', {'staff': staff})

        # 2. Validate new password match
        if new_password != confirm_password:
            messages.error(request, 'New password and confirm password do not match.')
            return render(request, 'account/expired_password.html', {'staff': staff})

        # 3. Validate new password length
        if len(new_password) < 6:
            messages.error(request, 'New password must be at least 6 characters long.')
            return render(request, 'account/expired_password.html', {'staff': staff})

        # 4. Prevent reusing the same password
        if old_password == new_password:
            messages.error(request, 'New password cannot be identical to your expired password.')
            return render(request, 'account/expired_password.html', {'staff': staff})

        # Update CompanyStaff
        staff.password = make_password(new_password)
        staff.password_changed_at = timezone.now()
        staff.is_authenticated = True
        staff.save()

        # Update Django User model if exists
        if staff.email:
            from account.models import User as AccountUser
            dj_user = AccountUser.objects.filter(email=staff.email).first()
            if dj_user:
                dj_user.set_password(new_password)
                dj_user.save()

        # Clear expired flag in session
        if 'password_expired' in request.session:
            del request.session['password_expired']
        request.session['company_staff_id'] = staff.id
        request.session["new_notification"] = getattr(staff, 'new_notification', False)

        sweetify.success(request, 'Password Updated Successfully', text='Your password has been renewed for the next 60 days.', timer=4000)
        messages.success(request, 'Password updated successfully!')

        # Route directly to user's dashboard based on role
        role = getattr(staff, 'role', None) or Login._role_from_flags(staff)
        is_hr_user = role == CompanyStaff.ROLE_HR or getattr(staff, 'is_hr', False)

        if role == CompanyStaff.ROLE_SUPERADMIN:
            return HttpResponseRedirect('/hrms/superadmin/')
        if (role == CompanyStaff.ROLE_ADMIN or staff.is_company_admin) and staff.company_id:
            return HttpResponseRedirect(f"/hrms/dashboard/admin/{staff.company_id}/{staff.pk}/")
        if is_hr_user and staff.company_id:
            return HttpResponseRedirect(f"/hrms/dashboard/hr/{staff.company_id}/{staff.pk}/")
        if (role == CompanyStaff.ROLE_MANAGER or staff.is_manager) and staff.company_id:
            return HttpResponseRedirect(f"/hrms/managers/dashboard/{staff.company_id}/{staff.pk}/")
        if (role == CompanyStaff.ROLE_EMPLOYEE or staff.is_employee) and staff.company_id:
            return HttpResponseRedirect(f"/hrms/employee/employee_dashboard/{staff.company_id}/{staff.pk}/")

        return HttpResponseRedirect('/hrms/')

    return render(request, 'account/expired_password.html', {'staff': staff})


def hr_dashboard(request, company_id, company_staff_id):
    company = get_object_or_404(Company, id=company_id)
    staff = get_object_or_404(CompanyStaff, id=company_staff_id, company=company)

    from django.db.models import Q
    from biometric.models import BiometricDevice, BiometricEventLog

    total_employees = Employee.objects.filter(Q(user__company=company) | Q(user__isnull=True)).count()
    total_managers = Manager.objects.filter(Q(user__company=company) | Q(user__isnull=True)).count()
    today_date = timezone.now().date()
    today_attendance = Attendance.objects.filter(
        check_in__date=today_date
    ).count()

    devices_count = BiometricDevice.objects.filter(company=company).count()
    recent_events = (
        BiometricEventLog.objects.filter(
            Q(company=company) | Q(device__company=company) | Q(company__isnull=True)
        )
        .exclude(biometric_user_id__icontains='fk_name')
        .exclude(biometric_user_id__icontains='{')
        .select_related('device', 'employee', 'manager')
        .order_by('-received_at', '-id')[:10]
    )

    # Personal HR Employee profile & Reporting Manager
    hr_employee = Employee.objects.filter(user=staff).first()
    if not hr_employee and staff.email:
        hr_employee = Employee.objects.filter(employee_email=staff.email).first()
    hr_manager = hr_employee.employee_reports_to if hr_employee else None
    hr_attendance = None
    hr_is_check_in = 'Check In'
    hr_hours_num = ''
    if hr_employee:
        today_start = timezone.now().replace(hour=0, minute=0, second=0, microsecond=0)
        today_end = timezone.now().replace(hour=23, minute=59, second=59, microsecond=999999)
        hr_attendance = Attendance.objects.filter(
            employee=hr_employee,
            check_in__gte=today_start,
            check_in__lte=today_end
        ).order_by('-id').first()

        if hr_attendance:
            if hr_attendance.check_out and hr_attendance.check_in:
                time_diff = hr_attendance.check_out - hr_attendance.check_in
                total_seconds = max(0, int(time_diff.total_seconds()))
                hours = total_seconds // 3600
                minutes = (total_seconds % 3600) // 60
                hr_hours_num = f"{hours}h {minutes}m"
                hr_is_check_in = 'Completed'
            elif hr_attendance.check_in and not hr_attendance.check_out:
                hr_is_check_in = 'Check Out'
                time_diff = timezone.now() - hr_attendance.check_in
                total_seconds = max(0, int(time_diff.total_seconds()))
                hours = total_seconds // 3600
                minutes = (total_seconds % 3600) // 60
                hr_hours_num = f"{hours}h {minutes}m"

    is_admin = bool(staff and (staff.is_company_admin or staff.role in [CompanyStaff.ROLE_ADMIN, CompanyStaff.ROLE_SUPERADMIN]))

    context = {
        'company': company,
        'staff': staff,
        'is_admin': is_admin,
        'company_id': company_id,
        'company_staff_id': company_staff_id,
        'company_staff_authenticated': True,
        'user_initials': 'ADM' if is_admin else 'HR',
        'total_employees': total_employees,
        'employee_count': total_employees,
        'total_managers': total_managers,
        'manager_count': total_managers,
        'today_attendance': today_attendance,
        'today_date': today_date,
        'devices_count': devices_count,
        'recent_events': recent_events,
        'employee': hr_employee,
        'company_staff_is_hr': True,
        'hr_employee': hr_employee,
        'hr_manager': hr_manager,
        'hr_attendance': hr_attendance,
        'hr_is_check_in': hr_is_check_in,
        'hr_hours_num': hr_hours_num,
    }
    return render(request, 'account/hr_dashboard.html', context)


def hr_biometric_monitor(request, company_id, company_staff_id):
    company = get_object_or_404(Company, id=company_id)
    staff = get_object_or_404(CompanyStaff, id=company_staff_id, company=company)

    from biometric.models import BiometricDevice, BiometricEventLog
    from django.db.models import Q

    devices = BiometricDevice.objects.filter(company=company).order_by('name', 'id')
    recent_events = (
        BiometricEventLog.objects.filter(
            Q(company=company) | Q(device__company=company) | Q(company__isnull=True),
            status__in=[BiometricEventLog.STATUS_APPLIED, BiometricEventLog.STATUS_UNMATCHED],
        )
        .exclude(biometric_user_id__icontains='fk_name')
        .exclude(biometric_user_id__icontains='{')
        .select_related('device', 'employee', 'manager')
        .order_by('-received_at', '-id')[:50]
    )

    context = {
        'company': company,
        'staff': staff,
        'company_id': company_id,
        'company_staff_id': company_staff_id,
        'devices': devices,
        'recent_events': recent_events,
    }
    return render(request, 'account/hr_biometric_monitor.html', context)


def hr_profile_view(request, company_id, company_staff_id):
    company = get_object_or_404(Company, id=company_id)
    staff = get_object_or_404(CompanyStaff, id=company_staff_id, company=company)

    from employee.models import Post

    # Auto-provision Employee profile for HR/Admin if missing
    is_staff_admin = bool(staff and (staff.is_company_admin or staff.role in [CompanyStaff.ROLE_ADMIN, CompanyStaff.ROLE_SUPERADMIN]))
    dept_name = 'Administration' if is_staff_admin else 'Human Resources'
    
    profile_dept, _ = Department.objects.get_or_create(
        company=company,
        department_name=dept_name
    )

    hr_employee = Employee.objects.filter(user=staff).first()
    if not hr_employee:
        reporting_manager = Manager.objects.filter(user__company=company).first() or Manager.objects.first()
        staff_id_num = staff.id
        prefix = "ADM" if is_staff_admin else "HR"
        emp_id_str = f"EIC/{prefix}/{2700 + staff_id_num}"
        bio_id_str = str(staff_id_num)
        if Employee.objects.filter(biometric_id=bio_id_str).exists():
            bio_id_str = f"BIO-{staff_id_num}"

        email_str = staff.email or ('Admin' if is_staff_admin else 'HR Manager')
        name_part = email_str.split('@')[0].replace('.', ' ').replace('_', ' ')
        parts = name_part.split()
        first_name = parts[0].capitalize() if parts else ('Admin' if is_staff_admin else 'HR')
        last_name = parts[1].capitalize() if len(parts) > 1 else ('Manager' if is_staff_admin else 'Manager')

        hr_employee = Employee.objects.create(
            user=staff,
            employee_first_name=first_name,
            employee_last_name=last_name,
            employee_email=staff.email or '',
            employee_joining_date=timezone.now().date(),
            employee_department=profile_dept,
            employee_designation='Company Administrator' if is_staff_admin else 'HR Manager',
            employee_id=emp_id_str,
            biometric_id=bio_id_str,
            employee_reports_to=reporting_manager,
            employee_status='Active'
        )
    elif is_staff_admin:
        updated_fields = []
        if hr_employee.employee_department and hr_employee.employee_department.department_name == 'Human Resources':
            hr_employee.employee_department = profile_dept
            updated_fields.append('employee_department')
        if hr_employee.employee_designation in ['HR Manager', 'HR Administrator', '']:
            hr_employee.employee_designation = 'Company Administrator'
            updated_fields.append('employee_designation')
        if hr_employee.employee_id and hr_employee.employee_id.startswith('EIC/HR/'):
            hr_employee.employee_id = hr_employee.employee_id.replace('EIC/HR/', 'EIC/ADM/')
            updated_fields.append('employee_id')
        if updated_fields:
            hr_employee.save(update_fields=updated_fields)

    if request.method == "POST":
        # Check for cropped image input (base64)
        cropped_image_data = request.POST.get('cropped-image-input', '')
        if cropped_image_data and cropped_image_data.startswith('data:image'):
            import base64
            import uuid
            from django.core.files.base import ContentFile
            from django.utils.text import slugify
            try:
                hr_employee.avatar_base64 = cropped_image_data
                format_part, imgstr = cropped_image_data.split(';base64,')
                ext = format_part.split('/')[-1] if '/' in format_part else 'jpg'
                name = f"hr_{slugify(hr_employee.employee_first_name or 'profile')}_{uuid.uuid4().hex[:8]}.{ext}"
                hr_employee.employee_profile_image = ContentFile(base64.b64decode(imgstr), name=name)
            except Exception as e:
                messages.error(request, f'Error processing profile picture: {str(e)}')
        elif 'employee_profile_image' in request.FILES:
            uploaded_image = request.FILES['employee_profile_image']
            hr_employee.employee_profile_image = uploaded_image
            try:
                import base64
                import mimetypes
                uploaded_image.seek(0)
                content = uploaded_image.read()
                uploaded_image.seek(0)
                content_type = mimetypes.guess_type(uploaded_image.name)[0] or 'image/jpeg'
                hr_employee.avatar_base64 = f"data:{content_type};base64,{base64.b64encode(content).decode('utf-8')}"
            except Exception:
                pass

        # Update editable fields
        fields_to_update = [
            'employee_first_name', 'employee_last_name',
            'employee_phone', 'employee_birth_date', 'employee_gender',
            'employee_address', 'employee_pin_code', 'employee_state', 'employee_country',
            'employee_marital_status', 'employee_emergency_primary_name',
            'employee_emergency_primary_relationship', 'employee_emergency_primary_phone1',
        ]
        for field in fields_to_update:
            val = request.POST.get(field)
            if val is not None:
                val_clean = val.strip()
                if field in ['employee_phone', 'employee_emergency_primary_phone1']:
                    digits = ''.join(filter(str.isdigit, val_clean.replace('+91', '')))
                    if len(digits) == 10:
                        val_clean = f"+91 {digits[:5]} {digits[5:]}"
                    elif digits:
                        val_clean = f"+91 {digits}"
                if field == 'employee_birth_date' and not val_clean:
                    continue
                setattr(hr_employee, field, val_clean)

        hr_employee.save()
        messages.success(request, 'HR Profile updated successfully!')
        return redirect('hr_profile', company_id=company_id, company_staff_id=company_staff_id)

    # Documents
    post = Post.objects.filter(user=hr_employee).first()

    is_admin = bool(staff and (staff.is_company_admin or staff.role in [CompanyStaff.ROLE_ADMIN, CompanyStaff.ROLE_SUPERADMIN]))

    context = {
        'company': company,
        'staff': staff,
        'is_admin': is_admin,
        'company_id': company_id,
        'company_staff_id': company_staff_id,
        'hr_employee': hr_employee,
        'employee': hr_employee,
        'company_staff_is_hr': True,
        'profile': hr_employee,
        'post': post,
    }
    return render(request, 'account/hr_profile.html', context)


def upload_hr_profile_image(request, company_id, company_staff_id):
    if request.method != 'POST':
        return JsonResponse({'error': 'Method not allowed'}, status=405)
    try:
        staff = CompanyStaff.objects.get(id=company_staff_id, company_id=company_id)
        emp = Employee.objects.filter(user=staff).first()
        if not emp:
            return JsonResponse({'error': 'HR employee profile not found'}, status=404)
        if 'employee_profile_image' not in request.FILES:
            return JsonResponse({'error': 'No image file provided'}, status=400)

        uploaded_image = request.FILES['employee_profile_image']
        emp.employee_profile_image = uploaded_image
        import base64
        import mimetypes
        uploaded_image.seek(0)
        content = uploaded_image.read()
        uploaded_image.seek(0)
        content_type = mimetypes.guess_type(uploaded_image.name)[0] or 'image/jpeg'
        emp.avatar_base64 = f"data:{content_type};base64,{base64.b64encode(content).decode('utf-8')}"
        emp.save()
        return JsonResponse({'success': True, 'image_url': emp.avatar_url})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


def remove_hr_profile_image(request, company_id, company_staff_id):
    if request.method != 'POST':
        return JsonResponse({'error': 'Method not allowed'}, status=405)
    try:
        staff = CompanyStaff.objects.get(id=company_staff_id, company_id=company_id)
        emp = Employee.objects.filter(user=staff).first()
        if not emp:
            return JsonResponse({'error': 'HR employee profile not found'}, status=404)
        emp.avatar_base64 = None
        if emp.employee_profile_image:
            emp.employee_profile_image.delete(save=False)
            emp.employee_profile_image = None
        emp.save()
        return JsonResponse({'success': True, 'fallback_url': emp.avatar_url})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


# ==============================================================================
# ATTENDANCE REGISTER MATRIX, DAILY ROSTER, YEARLY SUMMARY & EXPORT ENGINE
# ==============================================================================

def _get_attendance_register_data(company, period='monthly', date_str=None, month=None, year=None, dept_id=None, status_filter='all', search_query=None):
    import calendar
    from datetime import datetime, date, time
    from django.db.models import Q
    from administration.models import holiday
    from leave.models import Leave

    today = timezone.localdate()
    current_tz = timezone.get_current_timezone()

    # Parse Month & Year safely
    try:
        month = int(month) if month else today.month
        if month < 1 or month > 12:
            month = today.month
    except Exception:
        month = today.month

    try:
        year = int(year) if year else today.year
        if year < 2000 or year > 2100:
            year = today.year
    except Exception:
        year = today.year

    # Parse Target Date for Daily Mode
    selected_date = today
    if date_str:
        try:
            import dateutil.parser
            selected_date = dateutil.parser.parse(str(date_str)).date()
        except Exception:
            selected_date = today

    # Optimized Query: select_related department and user in 1 single SQL query
    emp_qs = Employee.objects.filter(
        Q(user__company=company) | Q(user__isnull=True)
    ).select_related('employee_department', 'user').order_by('employee_first_name', 'employee_last_name')

    if dept_id and str(dept_id).lower() != 'all':
        emp_qs = emp_qs.filter(employee_department_id=dept_id)

    if search_query:
        sq = search_query.strip()
        emp_qs = emp_qs.filter(
            Q(employee_first_name__icontains=sq) |
            Q(employee_last_name__icontains=sq) |
            Q(employee_id__icontains=sq) |
            Q(biometric_id__icontains=sq) |
            Q(employee_email__icontains=sq)
        )

    employees = list(emp_qs)
    departments = Department.objects.filter(company=company).order_by('department_name')

    # Pre-cache employee display metadata to eliminate repeated disk/property queries
    emp_meta = {}
    for emp in employees:
        is_female = bool(emp.employee_gender and str(emp.employee_gender).strip().lower() == 'female')
        if emp.avatar_base64 and emp.avatar_base64.strip():
            avatar_url = emp.avatar_base64.strip()
        elif emp.employee_image:
            avatar_url = f"/media/{emp.employee_image}"
        elif is_female:
            avatar_url = '/static/asets/images/dummy-woman.png'
        else:
            avatar_url = '/static/asets/images/dummy-man.png'

        desig_str = str(emp.employee_designation).strip() if emp.employee_designation else 'Staff'

        emp_meta[emp.id] = {
            'id': emp.id,
            'name': f"{emp.employee_first_name or ''} {emp.employee_last_name or ''}".strip() or (emp.employee_email or 'Staff'),
            'emp_id_display': emp.formatted_employee_id if hasattr(emp, 'formatted_employee_id') else (emp.employee_id or '-'),
            'bio_id': emp.biometric_id or '-',
            'dept_name': emp.employee_department.department_name if emp.employee_department else '-',
            'designation': desig_str,
            'avatar_url': avatar_url,
            'is_female': is_female,
            'obj': emp,
        }

    # Fetch official company holidays in 1 query
    holiday_dates = set()
    for h in holiday.objects.filter(company=company):
        try:
            import dateutil.parser
            h_dt = dateutil.parser.parse(str(h.date)).date()
            holiday_dates.add(h_dt)
        except Exception:
            pass

    # ==========================================================================
    # 1. DAILY VIEW (Single Day Detailed Punch Roster)
    # ==========================================================================
    if period == 'daily':
        day_start = timezone.make_aware(datetime(selected_date.year, selected_date.month, selected_date.day, 0, 0, 0), current_tz)
        day_end = timezone.make_aware(datetime(selected_date.year, selected_date.month, selected_date.day, 23, 59, 59, 999999), current_tz)

        is_saturday = selected_date.weekday() == 5
        is_sunday = selected_date.weekday() == 6
        is_weekend_day = is_saturday or is_sunday
        is_official_holiday = selected_date in holiday_dates
        is_non_working_day = is_weekend_day or is_official_holiday

        attendance_qs = Attendance.objects.filter(
            employee__in=employees,
            check_in__gte=day_start,
            check_in__lte=day_end
        ).order_by('check_in')

        daily_punch_map = {}
        for att in attendance_qs:
            emp_id = att.employee_id
            if emp_id not in daily_punch_map:
                daily_punch_map[emp_id] = []
            daily_punch_map[emp_id].append(att)

        leave_qs = Leave.objects.filter(
            user__in=employees,
            status__in=['approved', 'Approved'],
            startdate__lte=selected_date,
            enddate__gte=selected_date
        )
        on_leave_emp_ids = set(l.user_id for l in leave_qs)

        daily_rows = []
        kpi_present = 0
        kpi_absent = 0
        kpi_late = 0
        kpi_halfday = 0
        kpi_leave = 0

        for idx, emp in enumerate(employees, start=1):
            meta = emp_meta[emp.id]
            punches = daily_punch_map.get(emp.id, [])
            is_on_leave = emp.id in on_leave_emp_ids

            check_in_str = '--:--'
            check_out_str = '--:--'
            worked_str = '-'
            status_code = 'A'
            status_label = 'Absent'
            status_class = 'badge-absent'

            if punches:
                first_p = punches[0]
                last_p = punches[-1]
                ci_local = timezone.localtime(first_p.check_in)
                check_in_str = ci_local.strftime('%I:%M %p')

                if last_p.check_out:
                    co_local = timezone.localtime(last_p.check_out)
                    check_out_str = co_local.strftime('%I:%M %p')
                    diff_sec = int((co_local - ci_local).total_seconds())
                    if diff_sec > 0:
                        worked_str = f"{diff_sec // 3600}h {(diff_sec % 3600) // 60}m"
                elif first_p.check_in:
                    worked_str = "In Progress"

                is_late = ci_local.time() > time(10, 15)
                is_half_day = False
                if last_p.check_out:
                    diff_sec = int((timezone.localtime(last_p.check_out) - ci_local).total_seconds())
                    if 0 < diff_sec < 16200:
                        is_half_day = True

                if is_half_day:
                    status_code = 'HD'
                    status_label = 'Half Day'
                    status_class = 'badge-halfday'
                    kpi_halfday += 1
                    kpi_present += 1
                elif is_late:
                    status_code = 'L'
                    status_label = 'Late'
                    status_class = 'badge-late'
                    kpi_late += 1
                    kpi_present += 1
                else:
                    status_code = 'P'
                    status_label = 'Present'
                    status_class = 'badge-present'
                    kpi_present += 1
            elif is_on_leave:
                status_code = 'LV'
                status_label = 'On Leave'
                status_class = 'badge-leave'
                kpi_leave += 1
            elif is_non_working_day:
                status_code = 'H'
                status_label = 'Weekend Off (Sat/Sun)' if is_weekend_day else 'Holiday'
                status_class = 'badge-holiday'
            else:
                if selected_date > today:
                    status_code = '-'
                    status_label = 'Future Date'
                    status_class = 'badge-future'
                else:
                    status_code = 'A'
                    status_label = 'Absent'
                    status_class = 'badge-absent'
                    kpi_absent += 1

            include = True
            if status_filter == 'present' and status_code not in ['P', 'L', 'HD']:
                include = False
            elif status_filter == 'halfday' and status_code != 'HD':
                include = False
            elif status_filter == 'absent' and status_code != 'A':
                include = False
            elif status_filter == 'late' and status_code != 'L':
                include = False
            elif status_filter == 'leave' and status_code != 'LV':
                include = False

            if include:
                daily_rows.append({
                    'index': idx,
                    'employee': emp,
                    'meta': meta,
                    'check_in': check_in_str,
                    'check_out': check_out_str,
                    'worked': worked_str,
                    'status_code': status_code,
                    'status_label': status_label,
                    'status_class': status_class,
                    'punches_count': len(punches),
                })

        return {
            'period': 'daily',
            'selected_date': selected_date,
            'selected_date_str': selected_date.strftime('%Y-%m-%d'),
            'is_weekend_day': is_weekend_day,
            'is_official_holiday': is_official_holiday,
            'daily_rows': daily_rows,
            'departments': departments,
            'kpi': {
                'total_employees': len(employees),
                'filtered_employees': len(daily_rows),
                'total_records': len(attendance_qs),
                'present': kpi_present,
                'absent': kpi_absent,
                'late': kpi_late,
                'halfday': kpi_halfday,
                'leave': kpi_leave,
            }
        }

    # ==========================================================================
    # 2. YEARLY VIEW (12-Month Rollup Summary)
    # ==========================================================================
    elif period == 'yearly':
        yearly_start = timezone.make_aware(datetime(year, 1, 1, 0, 0, 0), current_tz)
        yearly_end = timezone.make_aware(datetime(year, 12, 31, 23, 59, 59, 999999), current_tz)

        attendance_qs = Attendance.objects.filter(
            employee__in=employees,
            check_in__gte=yearly_start,
            check_in__lte=yearly_end
        )

        emp_year_punches = {}
        for att in attendance_qs:
            emp_id = att.employee_id
            m = timezone.localtime(att.check_in).month
            d = timezone.localtime(att.check_in).day
            if emp_id not in emp_year_punches:
                emp_year_punches[emp_id] = {}
            if m not in emp_year_punches[emp_id]:
                emp_year_punches[emp_id][m] = set()
            emp_year_punches[emp_id][m].add(d)

        leave_qs = Leave.objects.filter(
            user__in=employees,
            status__in=['approved', 'Approved'],
            startdate__lte=yearly_end.date(),
            enddate__gte=yearly_start.date()
        )
        emp_year_leaves = {}
        for l in leave_qs:
            emp_id = l.user_id
            if emp_id not in emp_year_leaves:
                emp_year_leaves[emp_id] = {}
            if l.startdate and l.enddate:
                cur = max(l.startdate, yearly_start.date())
                fin = min(l.enddate, yearly_end.date())
                while cur <= fin:
                    m = cur.month
                    if m not in emp_year_leaves[emp_id]:
                        emp_year_leaves[emp_id][m] = 0
                    emp_year_leaves[emp_id][m] += 1
                    cur += timezone.timedelta(days=1)

        yearly_rows = []
        tot_present_annual = 0
        tot_absent_annual = 0
        tot_leaves_annual = 0

        for idx, emp in enumerate(employees, start=1):
            meta = emp_meta[emp.id]
            months_data = []
            ann_p = 0
            ann_a = 0
            ann_lv = 0
            ann_payable = 0

            for m in range(1, 13):
                m_days = calendar.monthrange(year, m)[1]
                m_punches_days = len(emp_year_punches.get(emp.id, {}).get(m, set()))
                m_leave_days = emp_year_leaves.get(emp.id, {}).get(m, 0)

                m_weekends = sum(1 for d_num in range(1, m_days + 1) if date(year, m, d_num).weekday() in [5, 6])
                m_holidays = sum(1 for d_num in range(1, m_days + 1) if date(year, m, d_num) in holiday_dates and date(year, m, d_num).weekday() not in [5, 6])
                total_off = m_weekends + m_holidays

                working_days = m_days - total_off
                m_absent = max(0, working_days - m_punches_days - m_leave_days)
                if date(year, m, 1) > today:
                    m_absent = 0

                m_payable = m_punches_days + total_off + m_leave_days

                ann_p += m_punches_days
                ann_a += m_absent
                ann_lv += m_leave_days
                ann_payable += m_payable

                months_data.append({
                    'month_num': m,
                    'month_abbr': calendar.month_abbr[m],
                    'present': m_punches_days,
                    'absent': m_absent,
                    'leaves': m_leave_days,
                    'payable': m_payable,
                    'is_future': date(year, m, 1) > today
                })

            tot_present_annual += ann_p
            tot_absent_annual += ann_a
            tot_leaves_annual += ann_lv

            include = True
            if status_filter == 'present' and ann_p == 0:
                include = False
            elif status_filter == 'absent' and ann_a == 0:
                include = False
            elif status_filter == 'leave' and ann_lv == 0:
                include = False

            if include:
                yearly_rows.append({
                    'index': idx,
                    'employee': emp,
                    'meta': meta,
                    'months': months_data,
                    'annual_present': ann_p,
                    'annual_absent': ann_a,
                    'annual_leaves': ann_lv,
                    'annual_payable': ann_payable,
                })

        return {
            'period': 'yearly',
            'current_year': year,
            'yearly_rows': yearly_rows,
            'departments': departments,
            'kpi': {
                'total_employees': len(employees),
                'filtered_employees': len(yearly_rows),
                'total_records': len(attendance_qs),
                'present': tot_present_annual,
                'absent': tot_absent_annual,
                'late': 0,
                'halfday': 0,
                'leave': tot_leaves_annual,
            }
        }

    # ==========================================================================
    # 3. MONTHLY VIEW (Full 1 to 31 Color-Coded Register Matrix)
    # ==========================================================================
    num_days = calendar.monthrange(year, month)[1]

    days_meta = []
    for d in range(1, num_days + 1):
        d_date = date(year, month, d)
        is_saturday = d_date.weekday() == 5
        is_sunday = d_date.weekday() == 6
        is_weekend = is_saturday or is_sunday
        is_official_hol = d_date in holiday_dates

        days_meta.append({
            'day': d,
            'date': d_date,
            'weekday_name': d_date.strftime('%a'),
            'is_saturday': is_saturday,
            'is_sunday': is_sunday,
            'is_weekend': is_weekend,
            'is_holiday': is_official_hol,
            'is_today': d_date == today,
            'is_future': d_date > today,
        })

    month_start = timezone.make_aware(datetime(year, month, 1, 0, 0, 0), current_tz)
    month_end = timezone.make_aware(datetime(year, month, num_days, 23, 59, 59, 999999), current_tz)

    attendance_qs = Attendance.objects.filter(
        employee__in=employees,
        check_in__gte=month_start,
        check_in__lte=month_end
    ).order_by('check_in')

    # Pre-aggregate punches by (emp_id, day) with formatted strings
    punch_summary_map = {}
    for att in attendance_qs:
        emp_id = att.employee_id
        ci_local = timezone.localtime(att.check_in)
        day_num = ci_local.day
        key = (emp_id, day_num)

        if key not in punch_summary_map:
            punch_summary_map[key] = {
                'first_ci': ci_local,
                'last_co': None,
                'check_in_str': ci_local.strftime('%I:%M %p'),
                'check_out_str': '--:--',
                'worked_str': 'In Progress',
                'is_late': ci_local.time() > time(10, 15),
                'is_half_day': False,
            }

        if att.check_out:
            co_local = timezone.localtime(att.check_out)
            punch_summary_map[key]['last_co'] = co_local
            punch_summary_map[key]['check_out_str'] = co_local.strftime('%I:%M %p')
            diff_sec = int((co_local - punch_summary_map[key]['first_ci']).total_seconds())
            if diff_sec > 0:
                punch_summary_map[key]['worked_str'] = f"{diff_sec // 3600}h {(diff_sec % 3600) // 60}m"
                if diff_sec < 16200:
                    punch_summary_map[key]['is_half_day'] = True

    leave_qs = Leave.objects.filter(
        user__in=employees,
        status__in=['approved', 'Approved'],
        startdate__lte=month_end.date(),
        enddate__gte=month_start.date()
    )
    leave_map = {}
    for l in leave_qs:
        emp_id = l.user_id
        if emp_id not in leave_map:
            leave_map[emp_id] = set()
        if l.startdate and l.enddate:
            cur_d = max(l.startdate, month_start.date())
            end_d = min(l.enddate, month_end.date())
            while cur_d <= end_d:
                leave_map[emp_id].add(cur_d)
                cur_d += timezone.timedelta(days=1)

    matrix_rows = []
    kpi_present = 0
    kpi_absent = 0
    kpi_late = 0
    kpi_halfday = 0
    kpi_leave = 0
    kpi_total_records = len(attendance_qs)

    for idx, emp in enumerate(employees, start=1):
        meta = emp_meta[emp.id]
        emp_days = []
        tot_p = 0
        tot_a = 0
        tot_l = 0
        tot_h = 0
        tot_lv = 0
        tot_hd = 0

        for dm in days_meta:
            d = dm['day']
            d_date = dm['date']
            is_future = dm['is_future']
            is_weekend = dm['is_weekend']
            is_holiday = is_weekend or dm['is_holiday']
            is_on_leave = d_date in leave_map.get(emp.id, set())
            punch_info = punch_summary_map.get((emp.id, d))

            cell_status = '-'
            cell_class = 'badge-future'
            cell_title = 'No Record'
            check_in_str = '--:--'
            check_out_str = '--:--'
            worked_str = '-'

            if is_future:
                cell_status = '-'
                cell_class = 'badge-future'
                cell_title = 'Future Date'
            elif punch_info:
                check_in_str = punch_info['check_in_str']
                check_out_str = punch_info['check_out_str']
                worked_str = punch_info['worked_str']

                if punch_info['is_half_day']:
                    cell_status = 'HD'
                    cell_class = 'badge-halfday'
                    cell_title = f"Half Day ({check_in_str} - {check_out_str})"
                    tot_hd += 1
                    tot_p += 1
                    kpi_halfday += 1
                    kpi_present += 1
                elif punch_info['is_late']:
                    cell_status = 'L'
                    cell_class = 'badge-late'
                    cell_title = f"Late ({check_in_str})"
                    tot_l += 1
                    tot_p += 1
                    kpi_late += 1
                    kpi_present += 1
                else:
                    cell_status = 'P'
                    cell_class = 'badge-present'
                    cell_title = f"Present ({check_in_str} - {check_out_str})"
                    tot_p += 1
                    kpi_present += 1
            elif is_on_leave:
                cell_status = 'LV'
                cell_class = 'badge-leave'
                cell_title = "Approved Leave"
                tot_lv += 1
                kpi_leave += 1
            elif is_holiday:
                cell_status = 'H'
                cell_class = 'badge-holiday'
                cell_title = "Weekend Off (Sat/Sun)" if is_weekend else "Official Holiday"
                tot_h += 1
            else:
                cell_status = 'A'
                cell_class = 'badge-absent'
                cell_title = "Absent"
                tot_a += 1
                kpi_absent += 1

            # When a specific status filter is active, only show matching day cells and blank out non-matching cells
            display_status = cell_status
            display_class = cell_class
            display_title = cell_title

            if status_filter == 'present':
                if cell_status not in ['P', 'L', 'HD']:
                    display_status = '-'
                    display_class = 'badge-future'
                    display_title = f"{cell_title} (Filtered out)"
            elif status_filter == 'halfday':
                if cell_status != 'HD':
                    display_status = '-'
                    display_class = 'badge-future'
                    display_title = f"{cell_title} (Filtered out)"
            elif status_filter == 'absent':
                if cell_status != 'A':
                    display_status = '-'
                    display_class = 'badge-future'
                    display_title = f"{cell_title} (Filtered out)"
            elif status_filter == 'late':
                if cell_status != 'L':
                    display_status = '-'
                    display_class = 'badge-future'
                    display_title = f"{cell_title} (Filtered out)"
            elif status_filter == 'leave':
                if cell_status != 'LV':
                    display_status = '-'
                    display_class = 'badge-future'
                    display_title = f"{cell_title} (Filtered out)"

            emp_days.append({
                'day': d,
                'status': display_status,
                'raw_status': cell_status,
                'class': display_class,
                'title': display_title,
                'check_in': check_in_str,
                'check_out': check_out_str,
                'worked': worked_str,
                'is_weekend': is_weekend,
                'is_saturday': dm['is_saturday'],
                'is_sunday': dm['is_sunday'],
                'is_today': dm['is_today'],
            })

        tot_payable = tot_p + tot_h + tot_lv

        include_emp = True
        if status_filter == 'present' and tot_p == 0:
            include_emp = False
        elif status_filter == 'halfday' and tot_hd == 0:
            include_emp = False
        elif status_filter == 'absent' and tot_a == 0:
            include_emp = False
        elif status_filter == 'late' and tot_l == 0:
            include_emp = False
        elif status_filter == 'leave' and tot_lv == 0:
            include_emp = False

        if include_emp:
            matrix_rows.append({
                'index': idx,
                'employee': emp,
                'meta': meta,
                'days': emp_days,
                'total_present': tot_p,
                'total_absent': tot_a,
                'total_late': tot_l,
                'total_holiday': tot_h,
                'total_leave': tot_lv,
                'total_halfday': tot_hd,
                'total_payable': tot_payable,
            })

    return {
        'period': 'monthly',
        'num_days': num_days,
        'days_meta': days_meta,
        'matrix_rows': matrix_rows,
        'departments': departments,
        'kpi': {
            'total_employees': len(employees),
            'filtered_employees': len(matrix_rows),
            'total_records': kpi_total_records,
            'present': kpi_present,
            'absent': kpi_absent,
            'late': kpi_late,
            'halfday': kpi_halfday,
            'leave': kpi_leave,
        }
    }


def monthly_attendance_register(request, company_id, company_staff_id):
    company = get_object_or_404(Company, id=company_id)
    staff = get_object_or_404(CompanyStaff, id=company_staff_id, company=company)

    import calendar
    today = timezone.localdate()

    period = request.GET.get('period', 'monthly')
    date_str = request.GET.get('date', today.strftime('%Y-%m-%d'))

    try:
        month = int(request.GET.get('month', today.month))
        if month < 1 or month > 12:
            month = today.month
    except Exception:
        month = today.month

    try:
        year = int(request.GET.get('year', today.year))
        if year < 2000 or year > 2100:
            year = today.year
    except Exception:
        year = today.year

    dept_id = request.GET.get('dept', 'all')
    status_filter = request.GET.get('status', 'all')
    search_query = request.GET.get('q', '').strip()

    data = _get_attendance_register_data(
        company=company,
        period=period,
        date_str=date_str,
        month=month,
        year=year,
        dept_id=dept_id,
        status_filter=status_filter,
        search_query=search_query
    )

    months_list = [(m, calendar.month_name[m]) for m in range(1, 13)]
    years_list = [y for y in range(today.year - 3, today.year + 2)]

    is_admin = bool(staff and (staff.is_company_admin or staff.role in [CompanyStaff.ROLE_ADMIN, CompanyStaff.ROLE_SUPERADMIN]))

    context = {
        'company': company,
        'staff': staff,
        'is_admin': is_admin,
        'company_id': company_id,
        'company_staff_id': company_staff_id,
        'period': period,
        'selected_date_str': data.get('selected_date_str', date_str),
        'selected_date': data.get('selected_date', today),
        'is_weekend_day': data.get('is_weekend_day', False),
        'current_month': month,
        'current_month_name': calendar.month_name[month],
        'current_year': year,
        'dept_id': dept_id,
        'status_filter': status_filter,
        'search_query': search_query,
        'months_list': months_list,
        'years_list': years_list,
        'days_meta': data.get('days_meta', []),
        'matrix_rows': data.get('matrix_rows', []),
        'daily_rows': data.get('daily_rows', []),
        'yearly_rows': data.get('yearly_rows', []),
        'departments': data.get('departments', []),
        'kpi': data.get('kpi', {}),
    }
    return render(request, 'account/attendance_register.html', context)


def export_attendance_register(request, company_id, company_staff_id):
    company = get_object_or_404(Company, id=company_id)
    staff = get_object_or_404(CompanyStaff, id=company_staff_id, company=company)

    import calendar
    today = timezone.localdate()

    period = request.GET.get('period', 'monthly')
    date_str = request.GET.get('date', today.strftime('%Y-%m-%d'))

    try:
        month = int(request.GET.get('month', today.month))
    except Exception:
        month = today.month

    try:
        year = int(request.GET.get('year', today.year))
    except Exception:
        year = today.year

    dept_id = request.GET.get('dept', 'all')
    status_filter = request.GET.get('status', 'all')
    search_query = request.GET.get('q', '').strip()
    export_format = request.GET.get('format', 'excel')

    data = _get_attendance_register_data(
        company=company,
        period=period,
        date_str=date_str,
        month=month,
        year=year,
        dept_id=dept_id,
        status_filter=status_filter,
        search_query=search_query
    )

    if period == 'daily':
        filename = f"Daily_Attendance_{data.get('selected_date_str', date_str)}"
    elif period == 'yearly':
        filename = f"Yearly_Attendance_Summary_{year}"
    else:
        month_name = calendar.month_name[month]
        filename = f"Attendance_Register_{month_name}_{year}"

    # CSV Export
    if export_format == 'csv':
        import csv
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        writer = csv.writer(response)

        if period == 'daily':
            writer.writerow(['#', 'Employee Name', 'Emp ID', 'Bio ID', 'Department', 'Punch In', 'Punch Out', 'Worked Hours', 'Status'])
            for row in data.get('daily_rows', []):
                emp = row['employee']
                emp_name = f"{emp.employee_first_name} {emp.employee_last_name}".strip()
                dept_name = emp.employee_department.department_name if emp.employee_department else '-'
                writer.writerow([
                    row['index'],
                    emp_name,
                    emp.formatted_employee_id if hasattr(emp, 'formatted_employee_id') else emp.employee_id,
                    emp.biometric_id or '-',
                    dept_name,
                    row['check_in'],
                    row['check_out'],
                    row['worked'],
                    row['status_label'],
                ])
        elif period == 'yearly':
            header = ['#', 'Employee Name', 'Emp ID', 'Bio ID', 'Department']
            for m in range(1, 13):
                header.append(calendar.month_abbr[m])
            header.extend(['Annual Present', 'Annual Absent', 'Annual Leaves', 'Total Annual Payable Days'])
            writer.writerow(header)
            for row in data.get('yearly_rows', []):
                emp = row['employee']
                emp_name = f"{emp.employee_first_name} {emp.employee_last_name}".strip()
                dept_name = emp.employee_department.department_name if emp.employee_department else '-'
                row_data = [
                    row['index'],
                    emp_name,
                    emp.formatted_employee_id if hasattr(emp, 'formatted_employee_id') else emp.employee_id,
                    emp.biometric_id or '-',
                    dept_name,
                ]
                for m_data in row['months']:
                    row_data.append(f"{m_data['present']} P / {m_data['payable']} Pay")
                row_data.extend([
                    row['annual_present'],
                    row['annual_absent'],
                    row['annual_leaves'],
                    row['annual_payable'],
                ])
                writer.writerow(row_data)
        else:
            header = ['#', 'Employee Name', 'Emp ID', 'Bio ID', 'Department']
            for dm in data.get('days_meta', []):
                tag = ' (Sat)' if dm['is_saturday'] else (' (Sun)' if dm['is_sunday'] else f" ({dm['weekday_name']})")
                header.append(f"{dm['day']}{tag}")
            header.extend(['Present (P)', 'Absent (A)', 'Late (L)', 'Half Day (HD)', 'Holiday/Weekend (H)', 'Leave (LV)', 'Total Payable Days'])
            writer.writerow(header)

            for row in data.get('matrix_rows', []):
                emp = row['employee']
                emp_name = f"{emp.employee_first_name} {emp.employee_last_name}".strip()
                dept_name = emp.employee_department.department_name if emp.employee_department else '-'
                row_data = [
                    row['index'],
                    emp_name,
                    emp.formatted_employee_id if hasattr(emp, 'formatted_employee_id') else emp.employee_id,
                    emp.biometric_id or '-',
                    dept_name,
                ]
                for d in row['days']:
                    row_data.append(d['status'])
                row_data.extend([
                    row['total_present'],
                    row['total_absent'],
                    row['total_late'],
                    row['total_halfday'],
                    row['total_holiday'],
                    row['total_leave'],
                    row['total_payable'],
                ])
                writer.writerow(row_data)
        return response

    # Excel .xlsx export
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from io import BytesIO

        wb = openpyxl.Workbook()
        ws = wb.active

        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )
        header_fill = PatternFill(start_color='1B2A4A', end_color='1B2A4A', fill_type='solid')
        header_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')

        fill_present = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
        fill_absent = PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')
        fill_late = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
        fill_halfday = PatternFill(start_color='FFE8D6', end_color='FFE8D6', fill_type='solid')
        fill_holiday = PatternFill(start_color='E8F0FE', end_color='E8F0FE', fill_type='solid')
        fill_leave = PatternFill(start_color='E2D9F3', end_color='E2D9F3', fill_type='solid')
        fill_future = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')

        if period == 'daily':
            ws.title = f"Daily_{data.get('selected_date_str')}"
            ws.merge_cells('A1:F1')
            ws['A1'] = f"DAILY ATTENDANCE ROSTER - {data.get('selected_date_str')}"
            ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
            ws['A1'].fill = header_fill

            headers = ['#', 'Employee Name', 'Emp ID', 'Bio ID', 'Department', 'Punch In', 'Punch Out', 'Worked Hours', 'Status']
            ws.append([])
            ws.append(headers)

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=3, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border

            curr_row = 4
            for row in data.get('daily_rows', []):
                emp = row['employee']
                emp_name = f"{emp.employee_first_name} {emp.employee_last_name}".strip()
                dept_name = emp.employee_department.department_name if emp.employee_department else '-'
                row_vals = [
                    row['index'],
                    emp_name,
                    emp.formatted_employee_id if hasattr(emp, 'formatted_employee_id') else emp.employee_id,
                    emp.biometric_id or '-',
                    dept_name,
                    row['check_in'],
                    row['check_out'],
                    row['worked'],
                    row['status_label']
                ]
                ws.append(row_vals)
                for col_idx in range(1, len(row_vals) + 1):
                    c = ws.cell(row=curr_row, column=col_idx)
                    c.border = thin_border
                    c.font = Font(name='Calibri', size=10)
                    c.alignment = Alignment(horizontal='center' if col_idx not in [2, 5] else 'left', vertical='center')
                    if col_idx == 9:
                        if row['status_code'] == 'HD':
                            c.fill = fill_halfday
                            c.font = Font(name='Calibri', size=10, bold=True, color='C2410C')
                        elif row['status_code'] == 'P':
                            c.fill = fill_present
                        elif row['status_code'] == 'A':
                            c.fill = fill_absent
                        elif row['status_code'] == 'L':
                            c.fill = fill_late
                        elif row['status_code'] == 'H':
                            c.fill = fill_holiday
                        elif row['status_code'] == 'LV':
                            c.fill = fill_leave
                curr_row += 1

        elif period == 'yearly':
            ws.title = f"Summary_{year}"
            ws.merge_cells('A1:G1')
            ws['A1'] = f"ANNUAL ATTENDANCE SUMMARY - {year}"
            ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
            ws['A1'].fill = header_fill

            headers = ['#', 'Employee Name', 'Emp ID', 'Bio ID', 'Department']
            for m in range(1, 13):
                headers.append(calendar.month_abbr[m])
            headers.extend(['Annual Present', 'Annual Absent', 'Annual Leaves', 'Total Payable Days'])
            ws.append([])
            ws.append(headers)

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=3, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border

            curr_row = 4
            for row in data.get('yearly_rows', []):
                emp = row['employee']
                emp_name = f"{emp.employee_first_name} {emp.employee_last_name}".strip()
                dept_name = emp.employee_department.department_name if emp.employee_department else '-'
                row_vals = [
                    row['index'],
                    emp_name,
                    emp.formatted_employee_id if hasattr(emp, 'formatted_employee_id') else emp.employee_id,
                    emp.biometric_id or '-',
                    dept_name,
                ]
                for m_data in row['months']:
                    row_vals.append(m_data['present'])
                row_vals.extend([
                    row['annual_present'],
                    row['annual_absent'],
                    row['annual_leaves'],
                    row['annual_payable']
                ])
                ws.append(row_vals)
                for col_idx in range(1, len(row_vals) + 1):
                    c = ws.cell(row=curr_row, column=col_idx)
                    c.border = thin_border
                    c.font = Font(name='Calibri', size=10)
                    c.alignment = Alignment(horizontal='center' if col_idx not in [2, 5] else 'left', vertical='center')
                curr_row += 1

        else: # Monthly Matrix
            month_name = calendar.month_name[month]
            ws.title = f"{month_name} {year}"
            ws.merge_cells('A1:G1')
            ws['A1'] = f"MONTHLY ATTENDANCE REGISTER (SAT & SUN OFF) - {month_name.upper()} {year}"
            ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
            ws['A1'].fill = header_fill

            headers = ['#', 'Employee Name', 'Emp ID', 'Bio ID', 'Department']
            for dm in data.get('days_meta', []):
                tag = '\nSat' if dm['is_saturday'] else ('\nSun' if dm['is_sunday'] else f"\n{dm['weekday_name']}")
                headers.append(f"{dm['day']}{tag}")
            headers.extend(['P', 'A', 'L', 'HD', 'H', 'LV', 'Payable Days'])

            ws.append([])
            ws.append(headers)

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=3, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = thin_border

            current_row = 4
            for row in data.get('matrix_rows', []):
                emp = row['employee']
                emp_name = f"{emp.employee_first_name} {emp.employee_last_name}".strip()
                dept_name = emp.employee_department.department_name if emp.employee_department else '-'
                emp_id_display = emp.formatted_employee_id if hasattr(emp, 'formatted_employee_id') else emp.employee_id

                row_cells = [
                    row['index'],
                    emp_name,
                    emp_id_display,
                    emp.biometric_id or '-',
                    dept_name,
                ]
                for d in row['days']:
                    row_cells.append(d['status'])
                row_cells.extend([
                    row['total_present'],
                    row['total_absent'],
                    row['total_late'],
                    row['total_halfday'],
                    row['total_holiday'],
                    row['total_leave'],
                    row['total_payable'],
                ])
                ws.append(row_cells)

                for col_idx, val in enumerate(row_cells, start=1):
                    c = ws.cell(row=current_row, column=col_idx)
                    c.border = thin_border
                    c.font = Font(name='Calibri', size=10)
                    if col_idx in [1, 3, 4]:
                        c.alignment = Alignment(horizontal='center', vertical='center')
                    elif col_idx == 2 or col_idx == 5:
                        c.alignment = Alignment(horizontal='left', vertical='center')
                    else:
                        c.alignment = Alignment(horizontal='center', vertical='center')

                    if 6 <= col_idx <= 5 + len(data.get('days_meta', [])):
                        if str(val) == 'P':
                            c.fill = fill_present
                            c.font = Font(name='Calibri', size=10, bold=True, color='155724')
                        elif str(val) == 'HD':
                            c.fill = fill_halfday
                            c.font = Font(name='Calibri', size=10, bold=True, color='C2410C')
                        elif str(val) == 'A':
                            c.fill = fill_absent
                            c.font = Font(name='Calibri', size=10, bold=True, color='721C24')
                        elif str(val) == 'L':
                            c.fill = fill_late
                            c.font = Font(name='Calibri', size=10, bold=True, color='856404')
                        elif str(val) == 'H':
                            c.fill = fill_holiday
                            c.font = Font(name='Calibri', size=10, bold=True, color='0C5460')
                        elif str(val) == 'LV':
                            c.fill = fill_leave
                            c.font = Font(name='Calibri', size=10, bold=True, color='381E72')
                        elif str(val) == '-':
                            c.fill = fill_future
                            c.font = Font(name='Calibri', size=10, color='888888')

                current_row += 1

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        return response
    except Exception as e:
        import csv
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        writer = csv.writer(response)
        writer.writerow(['Error generating Excel, fallback to CSV', str(e)])
        return response





